import tempfile
import unittest
from pathlib import Path

from docx import Document
from reportlab.pdfgen import canvas

from cv_review import (
    CVReviewError,
    extract_cv_text,
    format_cv_review,
    review_cv_file,
    review_cv_text,
)


class CVDocumentExtractionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.directory = Path(self.temp_dir.name)

    def make_pdf(self, name: str, lines: list[str]) -> Path:
        path = self.directory / name
        pdf = canvas.Canvas(str(path))
        y = 800
        for line in lines:
            pdf.drawString(60, y, line)
            y -= 24
        pdf.save()
        return path

    def make_docx(self, name: str, paragraphs: list[str]) -> Path:
        path = self.directory / name
        document = Document()
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
        document.save(path)
        return path

    def test_extracts_pdf_text_and_keeps_section_order(self) -> None:
        path = self.make_pdf(
            "candidate.pdf",
            ["Education", "Example College 2021", "Work Experience", "Assistant 2022"],
        )

        text = extract_cv_text(path)

        self.assertLess(text.index("Education"), text.index("Work Experience"))
        self.assertIn("Example College 2021", text)

    def test_extracts_docx_paragraphs_and_table_content_in_order(self) -> None:
        path = self.directory / "candidate.docx"
        document = Document()
        document.add_paragraph("Education")
        document.add_paragraph("Example College")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Work Experience"
        table.cell(0, 1).text = "Assistant"
        document.add_paragraph("Technical Skills")
        document.save(path)

        text = extract_cv_text(path)

        self.assertLess(text.index("Education"), text.index("Work Experience"))
        self.assertLess(text.index("Work Experience"), text.index("Technical Skills"))
        self.assertIn("Work Experience | Assistant", text)

    def test_rejects_unsupported_file_formats(self) -> None:
        path = self.directory / "candidate.txt"
        path.write_text("Education", encoding="utf-8")

        with self.assertRaisesRegex(CVReviewError, "Unsupported file type"):
            extract_cv_text(path)

    def test_rejects_empty_documents(self) -> None:
        path = self.directory / "empty.pdf"
        path.touch()

        with self.assertRaisesRegex(CVReviewError, "empty"):
            extract_cv_text(path)

    def test_reports_missing_and_unreadable_files(self) -> None:
        path = self.directory / "missing.docx"

        with self.assertRaisesRegex(CVReviewError, "could not be found"):
            extract_cv_text(path)

    def test_reports_pdf_extraction_failures_without_crashing(self) -> None:
        path = self.directory / "damaged.pdf"
        path.write_bytes(b"%PDF damaged document")

        with self.assertRaisesRegex(CVReviewError, "PDF could not be read"):
            extract_cv_text(path)

    def test_reports_docx_extraction_failures_without_crashing(self) -> None:
        path = self.directory / "damaged.docx"
        path.write_bytes(b"not a Word document")

        with self.assertRaisesRegex(CVReviewError, "DOCX file could not be opened"):
            extract_cv_text(path)

    def test_rejects_documents_without_extractable_text(self) -> None:
        path = self.make_pdf("image-only.pdf", [])

        with self.assertRaisesRegex(CVReviewError, "No readable text"):
            extract_cv_text(path)

    def test_formats_clear_review_sections_for_a_short_cv(self) -> None:
        text = (
            "Jordan Example\n"
            "jordan@example.com\n"
            "Education\n"
            "Example College\n"
            "Technical Skills\n"
            "Python\n"
        )

        report = format_cv_review("candidate.pdf", text)
        feedback = review_cv_text(text)

        for heading in (
            "What is already good:",
            "What needs improvement:",
            "Specific suggested corrections:",
            "Recommended next steps:",
        ):
            self.assertIn(heading, report)
        self.assertTrue(feedback["good"])
        self.assertTrue(feedback["improvements"])
        self.assertTrue(
            any(
                item.startswith("Work experience was not detected.")
                for item in feedback["improvements"]
            )
        )

    def test_flags_known_wording_without_inventing_replacements(self) -> None:
        feedback = review_cv_text(
            "Work Experience\nI was responsible of filing documents."
        )

        self.assertTrue(
            any("'responsible of' to 'responsible for'" in item
                for item in feedback["corrections"])
        )
        self.assertTrue(
            any("first-person wording" in item for item in feedback["corrections"])
        )
        self.assertTrue(
            any("verify" in item.casefold() or "accurate" in item.casefold()
                for item in feedback["next_steps"])
        )

    def test_review_file_works_end_to_end_for_docx(self) -> None:
        path = self.make_docx(
            "candidate.docx",
            ["Career Objective", "Seeking an entry-level role.", "Education", "Example College"],
        )

        report = review_cv_file(path)

        self.assertIn("candidate.docx", report)
        self.assertIn("Career Objective", report)
        self.assertIn("Education", report)
        self.assertIn("local rule-based checks", report)


if __name__ == "__main__":
    unittest.main()
